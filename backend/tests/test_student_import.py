"""Carga de estudiantes desde la plantilla y desde el listado oficial de curso."""

import io

import openpyxl
import pytest
from sqlalchemy import select

from conftest import ADMIN, COEDITOR, COORDINATOR, EVALUATOR, TestingSessionLocal, login
from app.models.entities import ECOEEvent, Station, Student
from app.utils.helpers import normalize_rut, rut_check_digit, rut_is_valid


def _rut_parts(body: int) -> tuple[int, str]:
    return body, rut_check_digit(str(body)).upper()


def official_list(count: int = 6, *, start_body: int = 21000001, extra_rows=()) -> bytes:
    """Listado oficial tal como lo entrega la universidad (encabezado + tabla)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A3"] = "LISTAS DE CURSOS"
    ws["A6"] = "UNID. ACADÉMICA : CARRERA MEDICINA"
    ws["A7"] = "AÑO      : 2026        SEMESTRE : Semestre 1"
    ws["A8"] = "ASIGNATURA    : MED1118 - CLINICA I"
    ws["A9"] = "SECCIÓN      : Seccion 1"
    ws["A11"] = "PROF TITULAR  : PROFESOR DE PRUEBA"
    ws.append([])
    ws.append([])
    ws.append(["CORRELATIVO", "RUT", "DV", "REGISTRO", "AP.PATERNO", "AP.MATERNO", "NOMBRES", "EMAIL"])
    for index in range(1, count + 1):
        body, dv = _rut_parts(start_body + index)
        dv_cell = int(dv) if dv.isdigit() else dv
        ws.append([index, body, dv_cell, 2026000000 + index, f"PATERNO{index}", f"MATERNO{index}",
                   f"NOMBRE {index}", f"alumno{index}@alu.uct.cl"])
    for row in extra_rows:
        ws.append(list(row))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def template_csv(rows: list[str], header: str = "nombre,apellidos,rut,correo,grupo,circuito") -> bytes:
    return ("\n".join([header, *rows]) + "\n").encode("utf-8")


@pytest.fixture
def empty_event():
    """ECOE en configuración, sin estudiantes ni estaciones."""
    with TestingSessionLocal() as db:
        base = db.get(ECOEEvent, 1)
        event = ECOEEvent(
            name="ECOE importación", date=base.date, course_name="MED", school_name="Salud",
            responsible_teacher="X", contact_email="x@ecoe.cl", circuit_mode="secuencial",
            station_time_minutes=5, transition_time_minutes=1, total_groups=1,
            status="en_configuracion",
        )
        db.add(event)
        db.commit()
        return event.id


def _add_circuits(event_id: int, names: list[str]) -> None:
    with TestingSessionLocal() as db:
        for number, name in enumerate(names, start=1):
            db.add(Station(
                ecoe_event_id=event_id, station_number=number, name=f"Estación {number}",
                station_type="procedimental", circuit_name=name, station_time_minutes=5,
                transition_time_minutes=1, expected_outcomes="x", student_activity="x",
                pre_entry_instruction="x", evaluator_instruction="x",
            ))
        db.commit()


def _import(client, event_id, data, filename="lista.xlsx", **params):
    query = "&".join(f"{key}={value}" for key, value in params.items())
    return client.post(
        f"/api/students/import?ecoe_event_id={event_id}{'&' + query if query else ''}",
        files={"file": (filename, data, "application/octet-stream")},
    )


def _students(event_id):
    with TestingSessionLocal() as db:
        return db.scalars(
            select(Student).where(Student.ecoe_event_id == event_id).order_by(Student.ecoe_number)
        ).all()


class TestRut:
    @pytest.mark.parametrize("raw,dv,expected", [
        ("12.345.678-5", None, "12345678-5"),
        ("12345678-5", None, "12345678-5"),
        (" 12345678 - 5 ", None, "12345678-5"),
        ("123456785", None, "12345678-5"),
        ("12345678", "5", "12345678-5"),
        (12345678, 5, "12345678-5"),
        ("6.000.000-K", None, "6000000-k") if rut_check_digit("6000000") == "k" else ("1-9", None, "1-9"),
        # Sin guion y el último dígito no calza: puede ser un cuerpo sin DV; no se parte.
        ("21000001", None, "21000001"),
        ("AB123456", None, "ab123456"),
        ("", None, ""),
    ])
    def test_canonical_form(self, raw, dv, expected):
        assert normalize_rut(raw, dv) == expected

    def test_check_digit(self):
        assert rut_is_valid("12345678-5")
        assert not rut_is_valid("12345678-9")
        assert not rut_is_valid("ab123456")
        assert not rut_is_valid("12345678")


class TestOfficialList:
    def test_preview_reads_the_official_list_and_writes_nothing(self, auth_client, empty_event):
        response = _import(auth_client, empty_event, official_list(6), dry_run="true")
        assert response.status_code == 200, response.text
        plan = response.json()
        assert plan["dry_run"] is True and plan["to_import"] == 6
        assert plan["source"]["format"] == "listado_oficial"
        assert plan["source"]["course"] == "MED1118 - CLINICA I"
        assert plan["source"]["section"] == "Seccion 1"
        assert plan["source"]["year"] == "2026"
        assert plan["source"]["semester"] == "Semestre 1"
        # Los profesores del encabezado no se leen.
        assert "PROFESOR" not in response.text
        first = plan["rows"][0]
        assert first["rut"] == normalize_rut(*_rut_parts(21000002))
        assert first["last_name"] == "PATERNO1 MATERNO1"
        assert first["name"] == "NOMBRE 1"
        assert first["registration_number"] == "2026000001"
        assert first["ecoe_number"] == "E001" and first["warnings"] == []
        assert plan["use_list_numbers"] is True
        assert _students(empty_event) == []

    def test_import_keeps_list_number_registration_and_canonical_rut(self, auth_client, empty_event):
        response = _import(auth_client, empty_event, official_list(6))
        assert response.status_code == 200 and response.json()["imported"] == 6
        students = _students(empty_event)
        assert [student.ecoe_number for student in students] == [f"E00{n}" for n in range(1, 7)]
        assert students[2].registration_number == "2026000003"
        assert all(rut_is_valid(student.rut) for student in students)
        assert {student.circuit_name for student in students} == {"Circuito A"}
        listed = auth_client.get(f"/api/students/{empty_event}").json()["items"]
        assert listed[0]["registration_number"]

    def test_reimport_does_not_duplicate(self, auth_client, empty_event):
        _import(auth_client, empty_event, official_list(6))
        again = _import(auth_client, empty_event, official_list(6)).json()
        assert again["imported"] == 0 and again["skipped_rut_duplicate"] == 6
        assert len(_students(empty_event)) == 6

    def test_same_rut_in_other_format_is_detected(self, auth_client, empty_event):
        body, dv = _rut_parts(21000002)
        dotted = f"{str(body)[:2]}.{str(body)[2:5]}.{str(body)[5:]}-{dv}"
        created = auth_client.post("/api/students", json={
            "ecoe_event_id": empty_event, "name": "Ya", "last_name": "Cargado", "rut": dotted,
            "email": "ya@alu.uct.cl", "group_name": "Grupo 1", "circuit_name": "Circuito A",
        })
        assert created.status_code == 200, created.text
        plan = _import(auth_client, empty_event, official_list(3), dry_run="true").json()
        assert [row["status"] for row in plan["rows"]] == ["ya_cargado", "nuevo", "nuevo"]
        # Ya hay estudiantes: no se pisan sus números, se continúa la numeración.
        assert plan["can_use_list_numbers"] is False
        assert [row["ecoe_number"] for row in plan["rows"]] == ["", "E002", "E003"]
        # Y el alta manual rechaza el mismo RUT escrito sin puntos.
        duplicate = auth_client.post("/api/students", json={
            "ecoe_event_id": empty_event, "name": "Otra", "last_name": "Vez", "rut": f"{body}-{dv}",
            "email": "otra@alu.uct.cl", "group_name": "Grupo 1", "circuit_name": "Circuito A",
        })
        assert duplicate.status_code == 400

    def test_second_section_continues_numbering(self, auth_client, empty_event):
        _import(auth_client, empty_event, official_list(4))
        second = _import(auth_client, empty_event, official_list(3, start_body=22000000)).json()
        assert second["imported"] == 3
        assert [student.ecoe_number for student in _students(empty_event)] == [f"E00{n}" for n in range(1, 8)]

    def test_problem_rows_are_reported_not_imported(self, auth_client, empty_event):
        body, dv = _rut_parts(21000002)
        extra = [
            [7, body, dv, 1, "REPETIDO", "X", "EN ARCHIVO", "rep@alu.uct.cl"],
            [8, "", "", 2, "SIN", "RUT", "ALGUIEN", "sinrut@alu.uct.cl"],
            [9, 23000001, rut_check_digit("23000001"), 3, "SIN", "CORREO", "ALGUIEN", ""],
            [10, 23000002, rut_check_digit("23000002"), 4, "CORREO", "MALO", "ALGUIEN", "no-es-correo"],
            [11, 23000003, "0" if rut_check_digit("23000003") != "0" else "1", 5, "DV", "MALO", "ALGUIEN", "dv@alu.uct.cl"],
        ]
        plan = _import(auth_client, empty_event, official_list(6, extra_rows=extra), dry_run="true").json()
        statuses = [row["status"] for row in plan["rows"][6:]]
        assert statuses == ["repetido", "sin_rut", "sin_correo", "sin_correo", "nuevo"]
        # DV que no calza: se carga, pero con aviso para revisar.
        assert plan["rows"][-1]["warnings"] and plan["with_warnings"] == 1
        result = _import(auth_client, empty_event, official_list(6, extra_rows=extra)).json()
        assert result["imported"] == 7
        assert result["skipped_rut_duplicate"] == 1 and result["skipped_missing_data"] == 3


class TestCircuitDistribution:
    def test_halves_by_list_order_when_event_has_mirror(self, auth_client, empty_event):
        _add_circuits(empty_event, ["Circuito A", "Circuito A", "Circuito B", "Circuito B"])
        plan = _import(auth_client, empty_event, official_list(7), dry_run="true").json()
        assert plan["distribution"] == "halves"
        assert plan["distribution_options"] == ["halves", "first"]
        assert plan["per_circuit"] == {"Circuito A": 4, "Circuito B": 3}
        assert [row["circuit_name"] for row in plan["rows"]] == ["Circuito A"] * 4 + ["Circuito B"] * 3
        _import(auth_client, empty_event, official_list(7))
        students = _students(empty_event)
        assert [s.circuit_name for s in students] == ["Circuito A"] * 4 + ["Circuito B"] * 3

    def test_can_force_everyone_to_first_circuit(self, auth_client, empty_event):
        _add_circuits(empty_event, ["Circuito A", "Circuito B"])
        plan = _import(auth_client, empty_event, official_list(4), dry_run="true", distribution="first").json()
        assert plan["per_circuit"] == {"Circuito A": 4}

    def test_single_circuit_event_offers_no_split(self, auth_client, empty_event):
        _add_circuits(empty_event, ["Piso 4"])
        plan = _import(auth_client, empty_event, official_list(4), dry_run="true", distribution="halves").json()
        assert plan["distribution"] == "first"
        assert plan["per_circuit"] == {"Piso 4": 4}

    def test_invalid_distribution_rejected(self, auth_client, empty_event):
        assert _import(auth_client, empty_event, official_list(2), distribution="azar").status_code == 422


class TestTemplate:
    def test_template_csv_still_works_and_respects_its_circuit_column(self, auth_client, empty_event):
        _add_circuits(empty_event, ["Circuito A", "Circuito B"])
        data = template_csv([
            "Ana,Soto Rey,12.345.678-5,ana@uct.cl,Grupo 2,circuito b",
            "Luis,Paz,11111111-1,luis@uct.cl,,",
            "Eva,Mar,PASAPORTE77,eva@uct.cl,,Circuito Z",
        ])
        plan = _import(auth_client, empty_event, data, "plantilla.csv", dry_run="true").json()
        assert plan["source"]["format"] == "plantilla"
        assert plan["distribution"] == "file"
        rows = plan["rows"]
        assert rows[0]["rut"] == "12345678-5" and rows[0]["group_name"] == "Grupo 2"
        assert rows[0]["circuit_name"] == "Circuito B"
        assert rows[1]["circuit_name"] == "Circuito A" and rows[1]["group_name"] == "Grupo 1"
        # Identificador que no es RUT: se acepta con aviso; circuito inexistente: aviso.
        assert rows[2]["status"] == "nuevo" and len(rows[2]["warnings"]) == 2
        assert plan["can_use_list_numbers"] is False
        assert _import(auth_client, empty_event, data, "plantilla.csv").json()["imported"] == 3

    def test_semicolon_latin1_csv(self, auth_client, empty_event):
        data = "nombre;apellidos;rut;correo\nJosé;Núñez;12345678-5;jose@uct.cl\n".encode("latin-1")
        plan = _import(auth_client, empty_event, data, "plantilla.csv", dry_run="true").json()
        assert plan["rows"][0]["name"] == "José" and plan["rows"][0]["last_name"] == "Núñez"

    def test_missing_columns_explained(self, auth_client, empty_event):
        response = _import(auth_client, empty_event, template_csv(["a,b"], header="uno,dos"), "x.csv")
        assert response.status_code == 400
        assert "uno" in response.json()["detail"]

    def test_not_a_spreadsheet(self, auth_client, empty_event):
        assert _import(auth_client, empty_event, b"basura", "lista.xlsx").status_code == 400


class TestPermissionsAndLocks:
    def test_coeditor_can_import(self, client):
        login(client, COEDITOR)
        response = _import(client, 1, official_list(2, start_body=24000000), dry_run="true")
        assert response.status_code == 200

    @pytest.mark.parametrize("credentials", [COORDINATOR, EVALUATOR])
    def test_other_roles_cannot_preview_or_import(self, client, empty_event, credentials):
        login(client, credentials)
        assert _import(client, empty_event, official_list(2), dry_run="true").status_code == 403
        assert _import(client, empty_event, official_list(2)).status_code == 403
        assert _students(empty_event) == []

    def test_anonymous_rejected(self, client, empty_event):
        assert _import(client, empty_event, official_list(2), dry_run="true").status_code == 401

    def test_closed_event_rejects_import(self, auth_client, empty_event):
        with TestingSessionLocal() as db:
            db.get(ECOEEvent, empty_event).status = "cerrado"
            db.commit()
        assert _import(auth_client, empty_event, official_list(2)).status_code == 409
        assert _students(empty_event) == []


def test_results_export_includes_registration(auth_client):
    with TestingSessionLocal() as db:
        student = db.scalars(select(Student).where(Student.ecoe_event_id == 1)).first()
        student.registration_number = "2026999999"
        db.commit()
    response = auth_client.get("/api/results/1/export/excel")
    assert response.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(response.content))["consolidado"]
    header = [cell.value for cell in sheet[1]]
    if sheet.max_row > 1:
        assert "matricula" in header
